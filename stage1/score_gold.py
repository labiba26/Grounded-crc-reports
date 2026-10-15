"""score_gold.py - scores the gold-standard labelling (THESIS_LOG.md, Phase 12).

Run it twice:

  1. After both annotators finish:
         python3 score_gold.py
     Prints agreement (percent and Cohen's kappa) per field and writes
     gold_labels/adjudication.xlsx - one row per disagreement, with both answers and
     both annotators' notes, and a yellow "Final" cell to fill in together.

  2. After filling in adjudication.xlsx, run it again. The gold standard is then every
     agreed answer plus every adjudicated one, and the script scores the extractor
     against it, and GDC too where GDC has the field. Results go to
     gold_labels/gold_scores.json.

Extractor accuracy is reported on the RANDOM stratum (30 cases, unbiased) and on
all 50; the enriched cases were picked for rare classes and would bias an overall
figure (see make_gold_sheet.py).

    python3 score_gold.py [--a PATH] [--b PATH]
"""
import argparse
import csv
import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = Path(__file__).resolve().parent
BASE = HERE.parent
GOLD = BASE / "gold_labels"
FIRST_ROW = 4          # row 3 is the worked example
FIELD_COLS = ["histologic_type", "mucinous_pct", "grade", "pt", "pn", "nodes_pos",
              "nodes_exam", "node_status", "lvi", "margins", "site", "post_treatment",
              "confidence", "notes"]
AGREE_FIELDS = ["histologic_type", "grade", "pt", "pn", "nodes_pos", "nodes_exam",
                "node_status", "lvi", "margins", "site", "post_treatment"]

# ---- gold answer -> the extractor's value space ---------------------------------
NONE = {"not stated", "form: answer unreadable", None, ""}


def g_type(v):
    if v in NONE: return None
    if v.startswith("mucinous adenocarcinoma"): return "mucinous"
    if v.startswith("signet"): return "signet_ring"
    if v == "other carcinoma": return "other_carcinoma"
    return "adenocarcinoma"            # NOS, and "with mucinous component" (WHO)


def g_grade(v):
    if v in NONE or "two-tier" in str(v): return None   # two-tier is not comparable
    return int(v[1])


def g_pt(v):
    if v in NONE: return None
    return "is" if v == "pTis" else int(v[2])


def g_pn(v):
    if v in NONE or v == "pNX": return None
    return int(v[2])


BOOL = {"positive": True, "negative": False, "present": True, "absent": False,
        "involved": True, "clear": False}
SITE = {"hepatic flexure": "hepatic_flexure", "splenic flexure": "splenic_flexure",
        "colon NOS": "colon_nos"}

COMPARE = [  # (name, gold field, gold->value, extractor field)
    ("histologic_type", "histologic_type", g_type, "histologic_type"),
    ("mucinous (WHO)", "histologic_type",
     lambda v: None if g_type(v) is None else g_type(v) in ("mucinous", "signet_ring"),
     "histologic_type"),
    ("grade", "grade", g_grade, "histologic_grade"),
    ("pt_stage", "pt", g_pt, "pt_stage"),
    ("pn_stage", "pn", g_pn, "pn_stage"),
    ("node_status", "node_status", lambda v: None if v in NONE else BOOL[v], "lymph_node_positive"),
    ("lymphovascular_invasion", "lvi", lambda v: None if v in NONE else BOOL[v], "lymphovascular_invasion"),
    ("margins_involved", "margins", lambda v: None if v in NONE else BOOL[v], "margins_involved"),
    ("tumour_site", "site", lambda v: None if v in NONE else SITE.get(v, v), "tumour_site"),
]


def norm(v):
    if v is None: return None
    if isinstance(v, float) and v.is_integer(): v = int(v)
    v = str(v).strip()
    return v or None


def read_sheet(path):
    ws = load_workbook(path, data_only=True)["Labels"]
    out = {}
    for row in ws.iter_rows(min_row=FIRST_ROW, values_only=True):
        if not row[0]:
            continue
        out[row[0]] = {"case_id": row[1], **{k: norm(v) for k, v in zip(FIELD_COLS, row[3:])}}
    return out, ws["B1"].value


def kappa(pairs):
    n = len(pairs)
    if not n: return None
    po = sum(a == b for a, b in pairs) / n
    ca, cb = Counter(a for a, _ in pairs), Counter(b for _, b in pairs)
    pe = sum(ca[c] * cb[c] for c in set(ca) | set(cb)) / n / n
    return None if pe >= 1 else (po - pe) / (1 - pe)


def write_adjudication(dis, path):
    wb = Workbook(); ws = wb.active; ws.title = "Adjudicate"
    head = ["Gold ID", "Case", "Field", "Annotator A", "Annotator B", "Notes A", "Notes B",
            "Final (fill in)"]
    for j, h in enumerate(head, 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
    for i, d in enumerate(dis, 2):
        for j, v in enumerate([d["gold_id"], d["case_id"], d["field"], d["a"], d["b"],
                               d["notes_a"], d["notes_b"], None], 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(name="Arial"); c.alignment = Alignment(wrap_text=True, vertical="top")
            if j == 8: c.fill = PatternFill("solid", fgColor="FFF2CC")
    for j, w in enumerate([8, 15, 16, 26, 26, 40, 40, 26], 1):
        ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = w
    ws.freeze_panes = "A2"
    ws["J1"] = "Type the agreed answer in Final, spelled exactly as in the drop-down lists " \
               "(or leave a number blank for 'not stated'). Then rerun score_gold.py."
    ws["J1"].font = Font(name="Arial", italic=True, color="7F7F7F")
    wb.save(path)


def read_adjudication(path):
    if not path.exists(): return None
    ws = load_workbook(path, data_only=True).active
    out, missing = {}, 0
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[0]: continue
        # a blank Final is a valid answer ("not stated") for the two count fields only
        if (r[7] is None or str(r[7]).strip() == "") and r[2] not in ("nodes_pos", "nodes_exam"):
            missing += 1
        out[(r[0], r[2])] = norm(r[7])
    return out, missing


def gdc_values():
    """GDC's T, N-status and mucinous per case, via compare_gdc's logic."""
    try:
        import compare_gdc as G
    except Exception:
        return {}
    rows = {r["submitter_id"]: r for r in csv.DictReader(open(G.GDC), delimiter="\t")}
    out = {}
    for cid, row in rows.items():
        d = G.colorectal_diagnosis(row)
        if d:
            out[cid] = {"pt_stage": G.gdc_t(d), "node_status": G.gdc_node_pos(d),
                        "mucinous (WHO)": G.gdc_mucinous(d)}
    return out


def score(gold, pred, ids):
    both = [(gold[i], pred[i]) for i in ids if gold[i] is not None and pred[i] is not None]
    res = {"n": len(ids),
           "gold_stated": sum(gold[i] is not None for i in ids),
           "extractor_stated": sum(pred[i] is not None for i in ids),
           "both_stated": len(both),
           "accuracy_when_both_stated": (sum(a == b for a, b in both) / len(both)) if both else None,
           "missed": sum(gold[i] is not None and pred[i] is None for i in ids),
           "invented": sum(gold[i] is None and pred[i] is not None for i in ids)}
    classes = sorted({str(a) for a, _ in both} | {str(b) for _, b in both})
    per = {}
    for c in classes:
        tp = sum(str(a) == c and str(b) == c for a, b in both)
        fp = sum(str(a) != c and str(b) == c for a, b in both)
        fn = sum(str(a) == c and str(b) != c for a, b in both)
        per[c] = {"gold_n": tp + fn, "precision": tp / (tp + fp) if tp + fp else None,
                  "recall": tp / (tp + fn) if tp + fn else None}
    res["per_class"] = per
    return res


def rel(p):
    try:
        return p.relative_to(BASE)
    except ValueError:
        return p


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--a", default=GOLD / "gold_labels_A.xlsx", type=Path)
    ap.add_argument("--b", default=GOLD / "gold_labels_B.xlsx", type=Path)
    ap.add_argument("--out", default=GOLD, type=Path)
    args = ap.parse_args()

    A, name_a = read_sheet(args.a); B, name_b = read_sheet(args.b)
    key = {r["gold_id"]: r for r in json.loads((GOLD / "sample_key.json").read_text())["rows"]}
    ids = sorted(key)
    blank_a = [(g, f) for g in ids for f in AGREE_FIELDS if f not in ("nodes_pos", "nodes_exam")
               and A[g][f] is None]
    blank_b = [(g, f) for g in ids for f in AGREE_FIELDS if f not in ("nodes_pos", "nodes_exam")
               and B[g][f] is None]
    print(f"annotator A: {name_a} · annotator B: {name_b}")
    if blank_a or blank_b:
        print(f"\nNOT FINISHED: {len(blank_a)} empty cells in A, {len(blank_b)} in B "
              f"(first few: {(blank_a + blank_b)[:5]}). Fill them, then rerun.")
        return 1

    # ---- 1. agreement ---------------------------------------------------------------
    print(f"\n=== Agreement between annotators ({len(ids)} reports) ===")
    agreement, dis = {}, []
    for f in AGREE_FIELDS:
        pairs = [(A[g][f], B[g][f]) for g in ids]
        agr = sum(a == b for a, b in pairs) / len(pairs)
        k = kappa(pairs)
        agreement[f] = {"agreement": agr, "kappa": k}
        print(f"  {f:16s} {agr:6.1%}   kappa {k:.2f}" if k is not None else f"  {f:16s} {agr:6.1%}")
        for g in ids:
            if A[g][f] != B[g][f]:
                dis.append({"gold_id": g, "case_id": key[g]["case_id"], "field": f,
                            "a": A[g][f], "b": B[g][f], "notes_a": A[g]["notes"],
                            "notes_b": B[g]["notes"]})
    adj_path = args.out / "adjudication.xlsx"
    adj = read_adjudication(adj_path)
    if adj is None:
        write_adjudication(dis, adj_path)
        print(f"\n{len(dis)} disagreements -> {rel(adj_path)}. Fill in the "
              f"'Final' column together, then rerun this script.")
        return 0
    finals, missing = adj
    if missing:
        print(f"\nadjudication.xlsx has {missing} empty 'Final' cells - finish those, then rerun.")
        return 1

    # ---- 2. the gold standard ------------------------------------------------------
    gold = {g: {f: (A[g][f] if A[g][f] == B[g][f] else finals.get((g, f))) for f in AGREE_FIELDS}
            for g in ids}
    (args.out / "gold_standard.json").write_text(json.dumps(
        {g: {"case_id": key[g]["case_id"], "stratum": key[g]["stratum"], **gold[g]} for g in ids},
        indent=1))

    # ---- 3. extractor (and GDC) against gold ----------------------------------------
    labels = json.loads((HERE / "structured_labels.json").read_text())
    gdc = gdc_values()
    rand = [g for g in ids if key[g]["stratum"] == "random"]
    results = {"annotators": [name_a, name_b], "agreement": agreement,
               "n_disagreements": len(dis), "extractor": {}, "gdc": {}}
    print("\n=== Extractor against the adjudicated gold standard ===")
    print(f"{'field':24s} {'random 30: acc (both stated)':>30s} {'missed':>7s} {'invented':>9s}"
          f" {'all 50: acc':>12s}")
    for name, gf, fn, ef in COMPARE:
        gv = {g: fn(gold[g][gf]) for g in ids}
        raw = {g: labels[key[g]["case_id"]]["labels"][ef] for g in ids}
        pv = {g: (None if raw[g] is None else
                  raw[g] in ("mucinous", "signet_ring") if name == "mucinous (WHO)" else raw[g])
              for g in ids}
        r_rand, r_all = score(gv, pv, rand), score(gv, pv, ids)
        results["extractor"][name] = {"random": r_rand, "all": r_all}
        fmt = lambda x: "  —  " if x is None else f"{x:.1%}"
        print(f"{name:24s} {fmt(r_rand['accuracy_when_both_stated']):>22s} "
              f"(n={r_rand['both_stated']:2d}) {r_rand['missed']:7d} {r_rand['invented']:9d}"
              f" {fmt(r_all['accuracy_when_both_stated']):>12s}")
        if name in ("pt_stage", "node_status", "mucinous (WHO)") and gdc:
            gd = {g: gdc.get(key[g]["case_id"], {}).get(name) for g in ids}
            results["gdc"][name] = {"all": score(gv, gd, ids)}
    if results["gdc"]:
        print("\n=== GDC against the gold standard (all 50) ===")
        for name, r in results["gdc"].items():
            a = r["all"]["accuracy_when_both_stated"]
            print(f"  {name:22s} {'—' if a is None else f'{a:.1%}'} (n={r['all']['both_stated']})")
    (args.out / "gold_scores.json").write_text(json.dumps(results, indent=1, default=str))
    print(f"\nwrote gold_standard.json and gold_scores.json in {rel(args.out)}/")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
