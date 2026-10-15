"""make_gold_sheet.py - builds the gold-standard labelling set (THESIS_LOG.md, Phase 12).

Two annotators label the same 50 reports independently, blind to the extractor and to
GDC. The adjudicated result measures extractor accuracy for every field, including
grade (which GDC does not carry), and the annotator agreement (Cohen's kappa)
shows how well-defined each field is.

SAMPLING - two strata, recorded in gold_labels/sample_key.json, never shown to annotators:
  * random   30 cases, uniform over all 590. Precision and coverage estimates come
             from this stratum alone, because it is unbiased.
  * enriched 20 cases from the rest, chosen for rare or hard cases the random draw
             would almost never contain (mucinous, G1, pT1, and reports where the
             extractor found nothing). They measure rare-class behaviour; they would
             bias an overall accuracy figure, so they are kept apart when scoring.
The 50 are shuffled together and numbered G01-G50, so an annotator cannot tell which
stratum a case came from.

    python3 make_gold_sheet.py      # writes gold_labels/ (refuses to overwrite filled sheets)
"""
import json
import random
import shutil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
BASE = HERE.parent
OUT = BASE / "gold_labels"
SEED = 2026
N_RANDOM = 30
ENRICH = [  # (label, predicate on extractor labels, how many)
    ("mucinous or signet-ring", lambda l: l["histologic_type"] in ("mucinous", "signet_ring"), 6),
    ("G1", lambda l: l["histologic_grade"] == 1, 5),
    ("pT1", lambda l: l["pt_stage"] == 1, 4),
    ("no histologic type found", lambda l: l["histologic_type"] is None, 2),
    ("no node status found", lambda l: l["lymph_node_positive"] is None, 3),
]

# ---- the fields, their allowed values, and the guide text for each ----------------
FORM = "form: answer unreadable"
NS = "not stated"
FIELDS = [
    ("histologic_type", "Histologic type", 30, [
        "adenocarcinoma NOS", "mucinous adenocarcinoma (>50%)",
        "adenocarcinoma with mucinous component (<=50% / partial / focal)",
        "signet-ring cell carcinoma (>50%)", "other carcinoma", NS, FORM]),
    ("mucinous_pct", "Mucinous % (if stated)", 12, None),
    ("grade", "Grade", 16, ["G1", "G2", "G3", "G4", "low grade (two-tier)",
                            "high grade (two-tier)", NS, FORM]),
    ("pt", "pT", 11, ["pT1", "pT2", "pT3", "pT4", "pTis", NS, FORM]),
    ("pn", "pN", 11, ["pN0", "pN1", "pN2", "pNX", NS, FORM]),
    ("nodes_pos", "Nodes positive (n)", 11, None),
    ("nodes_exam", "Nodes examined (n)", 11, None),
    ("node_status", "Node status", 14, ["positive", "negative", NS, FORM]),
    ("lvi", "Lymphovascular invasion", 14, ["present", "absent", NS, FORM]),
    ("margins", "Margins", 13, ["involved", "clear", NS, FORM]),
    ("site", "Tumour site", 16, ["cecum", "ascending", "hepatic flexure", "transverse",
                                 "splenic flexure", "descending", "sigmoid", "rectosigmoid",
                                 "rectum", "colon NOS", NS]),
    ("post_treatment", "Post-treatment (ypTNM)", 12, ["yes", "no"]),
    ("confidence", "Confidence", 11, ["sure", "unsure"]),
    ("notes", "Notes", 45, None),
]

GUIDE = """GOLD-STANDARD LABELLING GUIDE — TCGA colorectal pathology reports

WHAT YOU ARE DOING
Reading 50 pathology reports and recording what each one STATES, field by field.
Two people label the same 50 reports independently; disagreements are resolved
afterwards. Expect about 3 minutes per report (2–3 hours in total, in several sittings).

BEFORE YOU START
• Read the "Cheat sheet" tab: it translates every term you will meet into a drop-down option.
• Do the 5 practice reports first (practice/ folder, practice_sheet.xlsx), then check your
  answers against PRACTICE_ANSWERS.md. Practice cases are NOT among the 50.

KEEP IT BLIND
• Work alone. Do not discuss cases with the other annotator until both of you have finished.
• Do not open structured_labels.json, gdc_clinical_coad_read.tsv or the other annotator's sheet.
• Each report is a text file in the reports/ folder, named like its row (G01_…txt).
  The link in column C opens it.

HOW TO FILL IN THE SHEET
• Edit only the yellow cells in the "Labels" sheet. Most have a drop-down list — use it,
  so values are spelled identically. Numbers and notes are typed.
• Row 3 (grey) is a worked EXAMPLE (an invented case). Do not edit it.
• Put your name and the date at the top of the Labels sheet (yellow cells B1 and D1).
• Save often.

GENERAL RULES
1. Record what the report STATES, not what you infer or know. "Not stated" is not the
   same as negative — use it whenever the report does not say.
2. Printed forms: if the text lists a question with all its options (e.g.
   "Mucinous: ( ) No ( ) Yes ( ) Yes (Focal)") and you cannot see which was ticked,
   choose "form: answer unreadable". If the value is written after the label
   ("Grade: Moderately differentiated"), it is stated — record it.
3. Several specimens: label the main colorectal tumour. Ignore findings in other organs.
4. The report contradicts itself (e.g. "pN2" in one place, "0/16 nodes negative" in
   another): use the final diagnosis or synoptic summary, set Confidence to "unsure",
   and describe the conflict in Notes.
5. OCR errors are common: "pNO" means pN0; full stops appear mid-sentence.
6. Latin / Polish reports: bene differentiatum = G1 · mediocriter / medium = G2 ·
   male / parum differentiatum = G3 · "metastases carcinomatosae in lymphonodis" =
   node positive · "sine metastasibus" = negative · "partim mucinosum" = partially mucinous.
7. When unsure between two options, pick the better one, set Confidence to "unsure", and
   say why in Notes. Your notes are what makes adjudication fast.

FIELD BY FIELD
Histologic type — WHO rule: mucinous adenocarcinoma means >50% of the tumour is mucinous.
  • "Mucinous adenocarcinoma" with no percentage → mucinous adenocarcinoma (>50%).
  • "with mucinous features / component / differentiation", "partially / focally mucinous",
    or a stated percentage of 50% or less → adenocarcinoma with mucinous component.
  • Signet-ring cell carcinoma: the same >50% rule. "Signet-ring cell features" or
    "focal signet-ring cells" → adenocarcinoma NOS (mention it in Notes).
  • Medullary, neuroendocrine, squamous etc. → other carcinoma.
Mucinous % — only if a number is given. Leave blank otherwise.
Grade — G-number or words: well = G1, moderately = G2, poorly = G3, undifferentiated = G4.
  "Moderately to poorly" → the worse grade (G3), noted. Two-tier "low grade" /
  "high grade" → the two-tier options; do not convert them to G-numbers.
pT — only when a T category is written (pT3, T3, ypT3). T4a / T4b → pT4.
  If only the depth is described with no T code, choose "not stated" and copy the
  description into Notes.
pN — as written: N1a / N1b / N1c → pN1; N2a / N2b → pN2.
Nodes positive / examined — the two numbers, when given (e.g. "3/17" → 3 and 17).
Node status — positive: any regional lymph-node metastasis (a count of 1 or more, or
  pN1–pN2; for N1c, tumour deposits only, choose positive and write "N1c" in Notes).
  negative: nodes were examined and none were involved (0/n, pN0, "nodes negative").
  Distant metastasis does not count.
Lymphovascular invasion — present: lymphatic, vascular, venous or angiolymphatic
  invasion, or L1 / V1. absent: L0 and V0, or "no lymphovascular invasion".
  Perineural invasion is not lymphovascular invasion.
Margins — involved: tumour at any resection margin (R1, R2, "margin positive").
  clear: all margins free (R0, "margins negative"). A close but clear margin is clear.
Tumour site — the most specific site stated for the main tumour.
Post-treatment — yes if the report uses ypTNM or mentions neoadjuvant therapy,
  treatment effect or a tumour regression grade.
Confidence — sure / unsure, for the report as a whole.
"""


CHEAT = """CHEAT SHEET — what the words in these reports mean, and which option to pick

HOW TO USE THIS
Find the word or code you see in the report in the left column, pick the option on the
right. You never need to judge anything medically: the pathologist already decided, you
are only translating their words into the drop-down options.

HISTOLOGIC TYPE (what kind of tumour)
  adenocarcinoma · carcinoma · "adenocarcinoma of colorectal type"      → adenocarcinoma NOS
  tubular / tubulopapillary / papillary adenocarcinoma                   → adenocarcinoma NOS
  "adenocarcinoma arising in a (tubulo)villous adenoma"                  → adenocarcinoma NOS
  code 8140/3 (M8140/3, ICD-O 8140)                                      → adenocarcinoma NOS
  mucinous adenocarcinoma · mucinous carcinoma · colloid carcinoma      → mucinous adenocarcinoma (>50%)
  adenocarcinoma mucinosum · "tubulare et mucinosum" · code 8480/3      → mucinous adenocarcinoma (>50%)
  partially / focally mucinous · "partim mucinosum"                     → adenocarcinoma with mucinous component
  "with mucinous features / component / differentiation / areas"       → adenocarcinoma with mucinous component
  mucinous part given as 50% or less                                    → adenocarcinoma with mucinous component
  signet ring cell carcinoma · code 8490/3                              → signet-ring cell carcinoma (>50%)
  "signet ring cell features", "focal signet ring cells"               → adenocarcinoma NOS (+ note)
  medullary, neuroendocrine, squamous, adenosquamous, undifferentiated carcinoma → other carcinoma

GRADE (how abnormal the cells look)
  well differentiated · G1 · GI · "bene differentiatum"                 → G1
  moderately differentiated · G2 · GII · "medium / mediocriter differentiatum" → G2
  poorly differentiated · G3 · GIII · "male / parum differentiatum"     → G3
  undifferentiated · G4                                                  → G4
  "moderately to poorly differentiated" (two grades named)              → the worse one (G3) + note
  "low grade" / "high grade" (a two-level system)                        → low / high grade (two-tier)
  Careful: "high-grade dysplasia" describes a POLYP, not the cancer — ignore it.

pT — only when a T code is WRITTEN (pT1-4, T1-4, ypT…). T4a, T4b → pT4. Tis → pTis.
  For your understanding only (never use it to fill in pT without a code):
  T1 = into the submucosa · T2 = into the muscularis propria (muscle wall) ·
  T3 = through the muscle into the fat around the bowel (pericolic / mesorectal fat, subserosa) ·
  T4 = through the outer surface (serosa, peritoneum) or into another organ.

pN — as written. N0 = no nodes involved · N1 (N1a, N1b, N1c) = 1–3 nodes or tumour deposits ·
  N2 (N2a, N2b) = 4 or more nodes · NX = could not be assessed.
  OCR often turns zero into the letter O: "NO", "pNO" → pN0.
  Do NOT convert other staging systems: "Dukes B", "Astler-Coller B2", "Stage IIA" are not pN.

LYMPH NODE NUMBERS — positive first, examined second
  "3/19", "3 of 19", "three of nineteen"                                → 3 and 19
  "Number with metastasis: 0. Total number examined: 12"               → 0 and 12
  "(2 of 13 + 1)"                                                         → 2 and 13, note "+1"
  "0/16", "tumour-free lymph nodes (0/16)"                              → 0 and 16
  Latin / Polish: "(No XIV)" = number 14 (Roman numerals, see below).
  Only the number given → fill that one, leave the other blank.

NODE STATUS
  positive: a positive count of 1 or more · pN1 / pN2 · "metastatic carcinoma in 2 lymph nodes"
            · "metastases carcinomatosae in lymphonodis" · N1c (tumour deposits: + note "N1c")
  negative: 0/n · pN0 · "lymph nodes negative / free of tumour / without metastasis"
            · "lymphonodulitis reactiva" / "reactive lymph nodes" (inflamed, no cancer)
  not stated: nothing about regional lymph nodes.
  Distant metastasis (liver, lung, M1) is NOT node status — ignore it here.

LYMPHOVASCULAR INVASION (tumour inside small vessels)
  present: L1 · V1 · "lymphovascular / angiolymphatic / lymphatic / vascular / venous invasion
           present / identified"
  absent:  L0 and V0 (OCR: "LO", "VO") · "not identified" · "no lymphovascular invasion"
  Only one of L / V given: L1 or V1 alone → present; L0 or V0 alone → absent.
  IGNORE perineural invasion (Pn0 / Pn1) — it is not lymphovascular.

MARGINS (did the surgeon cut through tumour?)
  clear:    R0 · "margins free / negative / uninvolved / tumour-free" · "incision lines free"
            · a distance to the margin ("tumour 0.3 cm from the radial margin")
  involved: R1 · R2 · "margin involved / positive / infiltrated by tumour"
  Any margin counts (proximal, distal, radial, circumferential, mesenteric).

TUMOUR SITE
  cecum / caecum · ileocecal valve · Bauhin's valve · "valvula Bauhini"  → cecum
  ascending colon → ascending · hepatic flexure → hepatic flexure · transverse colon → transverse
  splenic flexure → splenic flexure · descending colon → descending · sigmoid (colon) → sigmoid
  rectosigmoid (junction) → rectosigmoid · rectum, rectal, "recti" → rectum
  only "colon" / "large intestine" / "right / left hemicolectomy" with no part named → colon NOS

POST-TREATMENT = yes if: ypT / ypN · "neoadjuvant" · "after chemo(radio)therapy" ·
  "treatment effect" · "tumour regression grade (TRG)". Otherwise → no.

OCR AND FORMAT TRAPS
  O instead of 0 (NO, pNO, LO, VO) · I instead of 1 (GI, GII) · full stops in the middle of
  sentences · text from a printed form: "( ) No ( ) Yes", "r No r Yes", "C No C Yes" — if you
  cannot see which box was ticked → "form: answer unreadable".
  A value written AFTER a label on a form IS stated: "Grade : Moderately Differentiated" → G2.

ROMAN NUMERALS (Polish / Latin reports)
  I 1 · II 2 · III 3 · IV 4 · V 5 · VI 6 · VII 7 · VIII 8 · IX 9 · X 10 · XI 11 · XII 12 ·
  XIII 13 · XIV 14 · XV 15 · XVI 16 · XVII 17 · XVIII 18 · XIX 19 · XX 20

THINGS YOU CAN IGNORE
  Dukes / Astler-Coller stages · M stage · tumour size · tumour budding · TILs / lymphocytes ·
  MSI / mismatch repair / immunostains · polyps elsewhere in the bowel · the gross
  (macroscopic) description when a final diagnosis exists.
"""

# Five practice reports (not in the 50), with worked answers - for training only.
PRACTICE = [
 ("P1", "TCGA-CM-6170", "a structured US report ('synoptic': Label: value)", {
   "histologic_type": ("adenocarcinoma NOS", "'Tumor Type: Adenocarcinoma'."),
   "mucinous_pct": (None, "No percentage given - leave blank."),
   "grade": ("G2", "'Histologic Grade: Moderately differentiated'."),
   "pt": ("pT2", "'Tumor Staging ... : pT2'."),
   "pn": ("pN0", "'Lymph Node Stage ... : NO' - OCR has turned the zero into the letter O."),
   "nodes_pos": (0, "'Number with metastasis: 0'."),
   "nodes_exam": (12, "'Total number examined: 12'."),
   "node_status": ("negative", "Nodes were examined and none had cancer."),
   "lvi": ("absent", "'Lymphovascular Invasion: Not identified'. The perineural line is ignored."),
   "margins": ("clear", "'Surgical Margins: Free of tumor'."),
   "site": ("descending", "'Tumor Location: Descending colon'."),
   "post_treatment": ("no", "Nothing about treatment before surgery."),
   "confidence": ("sure", "")}),
 ("P2", "TCGA-AA-3860", "a short German-style report (everything in one coded line)", {
   "histologic_type": ("adenocarcinoma NOS", "It says 'colon carcinoma', and the code M8140/3 means adenocarcinoma."),
   "mucinous_pct": (None, ""),
   "grade": ("G3", "'poorly differentiated ... G3'."),
   "pt": ("pT3", "'pT3'."),
   "pn": ("pN1", "'pN1'."),
   "nodes_pos": (2, "'pN1 (2 of 13 + 1)': 2 positive of 13."),
   "nodes_exam": (13, "The '+1' is unexplained (probably a tumour deposit) - write it in Notes."),
   "node_status": ("positive", "2 nodes positive."),
   "lvi": ("present", "'L1, V1'."),
   "margins": ("clear", "'free resection margins', 'R0'."),
   "site": ("sigmoid", "'sigmoid colon carcinoma'."),
   "post_treatment": ("no", ""),
   "confidence": ("sure", "Notes: '+1 after the node count, unclear'.")}),
 ("P3", "TCGA-D5-6924", "a Polish / Latin report", {
   "histologic_type": ("adenocarcinoma NOS", "'Adenocarcinoma tubulopapillare' is a growth pattern of ordinary adenocarcinoma."),
   "mucinous_pct": (None, ""),
   "grade": ("G1", "'(GI)' - the letter I is the Roman numeral 1."),
   "pt": ("pT3", "'pT3' in the final line."),
   "pn": ("pN0", "'pNO' - letter O for zero."),
   "nodes_pos": (0, "'Lymphonodulitis reactiva' = reactive (inflamed) nodes, no cancer."),
   "nodes_exam": (8, "'(NO VIII)' = number VIII = 8 nodes."),
   "node_status": ("negative", "pN0, reactive nodes only."),
   "lvi": ("not stated", "No L, V or vessel invasion mentioned."),
   "margins": ("clear", "'Incision lines free of neoplastic lesions'."),
   "site": ("sigmoid", "'sigmoid colon'."),
   "post_treatment": ("no", ""),
   "confidence": ("sure", "'Dukes B, Astler-Coller B2' are other staging systems - ignore them.")}),
 ("P4", "TCGA-DM-A1D9", "a scanned tick-box form", {
   "histologic_type": ("adenocarcinoma NOS", "'Histology : Adenocarcinoma' - the value is written after the label, so it is stated. The mucinous tick box itself is unreadable: say so in Notes."),
   "mucinous_pct": (None, ""),
   "grade": ("G2", "'Grade : Moderately Differentiated' - written after the label, so stated."),
   "pt": ("not stated", "No T code anywhere."),
   "pn": ("not stated", ""),
   "nodes_pos": (None, ""),
   "nodes_exam": (None, ""),
   "node_status": ("not stated", "Nothing about lymph nodes."),
   "lvi": ("form: answer unreadable", "'Angiolymphatic Invasion : r. No r Yes r Unknown' - all options printed, no visible tick."),
   "margins": ("not stated", ""),
   "site": ("cecum", "'Tumor Site: Cecum'."),
   "post_treatment": ("no", ""),
   "confidence": ("unsure", "Scanned form with garbled text - 'unsure' is right here.")}),
 ("P5", "TCGA-AA-3949", "a 'partially mucinous' tumour", {
   "histologic_type": ("adenocarcinoma with mucinous component (<=50% / partial / focal)", "'partially mucinous adenocarcinoma' - partial means not >50%."),
   "mucinous_pct": (None, "No number given."),
   "grade": ("G3", "'poorly differentiated', and 'G3' at the end."),
   "pt": ("pT3", "'pT3'."),
   "pn": ("pN1", "'pN1 (3/19)'."),
   "nodes_pos": (3, "'Three of nineteen', '3/19'."),
   "nodes_exam": (19, ""),
   "node_status": ("positive", "3 nodes have metastasis."),
   "lvi": ("absent", "'LO V0' - L0 with a letter O, and V0."),
   "margins": ("clear", "'resection margins are tumor-free', 'R0'."),
   "site": ("cecum", "'0.5 cm aborally to Bauhin's valve' - Bauhin's valve is the ileocecal valve, i.e. the cecum (see the cheat sheet)."),
   "post_treatment": ("no", ""),
   "confidence": ("sure", "")}),
]


def load_text(case):
    parts = [(BASE / p).read_text(errors="ignore") for p in case["reports"]]
    return ("\n\n----- next report document -----\n\n").join(parts)


def sample():
    labels = json.loads((HERE / "structured_labels.json").read_text())
    index = {c["case_id"]: c for c in json.loads(
        (BASE / "dataset_manifests_final" / "dataset_index.json").read_text())}
    split = {}
    for s in ("train_coad", "test_coad", "holdout_read"):
        for c in json.loads((BASE / "dataset_manifests_final" / f"{s}.json").read_text()):
            split[c["case_id"]] = s
    rng = random.Random(SEED)
    ids = sorted(index)
    chosen = [(c, "random", "") for c in rng.sample(ids, N_RANDOM)]
    taken = {c for c, _, _ in chosen}
    for why, pred, k in ENRICH:
        pool = [c for c in ids if c not in taken and pred(labels[c]["labels"])]
        for c in rng.sample(pool, min(k, len(pool))):
            chosen.append((c, "enriched", why)); taken.add(c)
    rng.shuffle(chosen)
    return [{"gold_id": f"G{i:02d}", "case_id": c, "stratum": st, "enriched_for": why,
             "split": split.get(c), "cohort": index[c]["cohort"]}
            for i, (c, st, why) in enumerate(chosen, 1)], index


ARIAL = "Arial"
YELLOW = PatternFill("solid", fgColor="FFF2CC")
GREY = PatternFill("solid", fgColor="EDEDED")
HEAD = PatternFill("solid", fgColor="1F3864")
THIN = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
EXAMPLE = {"histologic_type": "adenocarcinoma with mucinous component (<=50% / partial / focal)",
           "mucinous_pct": 20, "grade": "G2", "pt": "pT3", "pn": "pN1", "nodes_pos": 2,
           "nodes_exam": 17, "node_status": "positive", "lvi": "present", "margins": "clear",
           "site": "sigmoid", "post_treatment": "no", "confidence": "sure",
           "notes": "Invented example. 'Mucinous component approx. 20%' -> not mucinous "
                    "adenocarcinoma. Nodes 2/17, pN1b."}


def build_workbook(rows, annotator, folder="reports"):
    wb = Workbook()

    # ---- Guide ------------------------------------------------------------------
    g = wb.active; g.title = "Guide"
    g.column_dimensions["A"].width = 110
    for i, line in enumerate(GUIDE.strip("\n").split("\n"), 1):
        c = g.cell(row=i, column=1, value=line)
        is_head = bool(line) and line.split(" —")[0].isupper()
        c.font = Font(name=ARIAL, size=13 if i == 1 else 11, bold=is_head,
                      color="1F3864" if is_head else "000000")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    r = len(GUIDE.strip("\n").split("\n")) + 2
    g.cell(row=r, column=1, value="LEGEND").font = Font(name=ARIAL, bold=True, color="1F3864")
    g.cell(row=r + 1, column=1, value="Yellow cell = fill this in").fill = YELLOW
    g.cell(row=r + 2, column=1, value="Grey row = worked example, do not edit").fill = GREY
    for k in (1, 2):
        g.cell(row=r + k, column=1).font = Font(name=ARIAL)

    # ---- Cheat sheet ----------------------------------------------------------------
    cs = wb.create_sheet("Cheat sheet")
    cs.column_dimensions["A"].width = 120
    for i, line in enumerate(CHEAT.strip("\n").split("\n"), 1):
        c = cs.cell(row=i, column=1, value=line)
        head = bool(line) and not line.startswith(" ") and line.split(" —")[0].split(" (")[0].isupper()
        c.font = Font(name=ARIAL, size=13 if i == 1 else 10, bold=head,
                      color="1F3864" if head else "000000")
        c.alignment = Alignment(wrap_text=True, vertical="top")

    # ---- Lists (hidden) for drop-downs --------------------------------------------
    lst = wb.create_sheet("Lists")
    col = 1
    ranges = {}
    for key, _, _, options in FIELDS:
        if not options:
            continue
        lst.cell(row=1, column=col, value=key).font = Font(name=ARIAL, bold=True)
        for j, o in enumerate(options, 2):
            lst.cell(row=j, column=col, value=o).font = Font(name=ARIAL)
        L = lst.cell(row=1, column=col).column_letter
        ranges[key] = f"Lists!${L}$2:${L}${len(options) + 1}"
        col += 1
    lst.sheet_state = "hidden"

    # ---- Labels -------------------------------------------------------------------
    ws = wb.create_sheet("Labels", 1)
    ws["A1"] = "Annotator:"; ws["B1"] = annotator if annotator != "B" else None
    ws["C1"] = "Date:"
    for ref in ("B1", "D1"):
        ws[ref].fill = YELLOW; ws[ref].border = THIN
    ws["B1"].comment = Comment("Your name", "gold sheet")
    ws["D1"].comment = Comment("Date you finished", "gold sheet")
    for ref in ("A1", "B1", "C1", "D1"):
        ws[ref].font = Font(name=ARIAL, bold=ref in ("A1", "C1"))
    ws["F1"] = "Edit only yellow cells. Row 3 is a worked example. Guide and Cheat sheet tabs have the rules."
    ws["F1"].font = Font(name=ARIAL, italic=True, color="7F7F7F")

    head = ["Gold ID", "Case", "Report file"] + [h for _, h, _, _ in FIELDS]
    widths = [8, 17, 27] + [w for _, _, w, _ in FIELDS]
    HR = 2
    for j, (h, w) in enumerate(zip(head, widths), 1):
        c = ws.cell(row=HR, column=j, value=h)
        c.font = Font(name=ARIAL, bold=True, color="FFFFFF"); c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = THIN
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[HR].height = 34

    first = 3
    ex = ["EXAMPLE", "(invented)", "—"] + [EXAMPLE.get(k) for k, _, _, _ in FIELDS]
    for j, v in enumerate(ex, 1):
        c = ws.cell(row=first, column=j, value=v)
        c.font = Font(name=ARIAL, italic=True, color="595959"); c.fill = GREY; c.border = THIN
        c.alignment = Alignment(wrap_text=True, vertical="top")

    for i, row in enumerate(rows, first + 1):
        ws.cell(row=i, column=1, value=row["gold_id"])
        ws.cell(row=i, column=2, value=row["case_id"])
        fname = f"{row['gold_id']}_{row['case_id']}.txt"
        c = ws.cell(row=i, column=3, value=fname)
        c.hyperlink = f"{folder}/{fname}"
        c.font = Font(name=ARIAL, color="0563C1", underline="single")
        for j in (1, 2):
            ws.cell(row=i, column=j).font = Font(name=ARIAL, bold=j == 1)
        for j in range(1, len(head) + 1):
            c = ws.cell(row=i, column=j); c.border = THIN
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if j > 3:
                c.fill = YELLOW; c.font = Font(name=ARIAL)
    last = first + len(rows)

    for j, (key, _, _, options) in enumerate(FIELDS, 4):
        L = ws.cell(row=HR, column=j).column_letter
        if options:
            dv = DataValidation(type="list", formula1=ranges[key], allow_blank=True,
                                showErrorMessage=True, errorTitle="Use the list",
                                error="Pick a value from the drop-down (see the Guide tab).")
        else:
            if key == "notes":
                continue
            dv = DataValidation(type="whole", operator="between", formula1="0",
                                formula2="200" if key != "mucinous_pct" else "100",
                                allow_blank=True, showErrorMessage=True,
                                error="Whole number only, or leave blank.")
        ws.add_data_validation(dv)
        dv.add(f"{L}{first + 1}:{L}{last}")
    ws.freeze_panes = ws.cell(row=first + 1, column=4)
    wb.active = 1
    return wb


def build_practice(index):
    """practice/ - 5 reports not in the 50, a sheet, and the worked answers."""
    pdir = OUT / "practice"; pdir.mkdir(exist_ok=True)
    gold_ids = {r["case_id"] for r in sample()[0]}
    rows = []
    for pid, cid, _, _ in PRACTICE:
        assert cid not in gold_ids, f"practice case {cid} is in the gold set"
        (pdir / f"{pid}_{cid}.txt").write_text(load_text(index[cid]))
        rows.append({"gold_id": pid, "case_id": cid})
    build_workbook(rows, None, folder="practice").save(OUT / "practice_sheet.xlsx")
    names = {k: h for k, h, _, _ in FIELDS}
    md = ["# Practice answers", "",
          "Check your practice sheet against these **after** you have filled it in. "
          "Every answer says which words in the report it comes from.", ""]
    for pid, cid, kind, ans in PRACTICE:
        md += [f"## {pid} — {cid} — {kind}", "", "| Field | Answer | Why |", "|---|---|---|"]
        for k, (v, why) in ans.items():
            md.append(f"| {names[k]} | {'(blank)' if v is None else v} | {why} |")
        md.append("")
    (OUT / "PRACTICE_ANSWERS.md").write_text("\n".join(md))


def build():
    rows, index = sample()
    OUT.mkdir(exist_ok=True)
    for name in ("gold_labels_A.xlsx", "gold_labels_B.xlsx"):
        if (OUT / name).exists():
            raise SystemExit(f"{OUT / name} exists - refusing to overwrite a sheet that "
                             f"may already be filled in. Move it aside first.")
    rep = OUT / "reports"; rep.mkdir(exist_ok=True)
    for r in rows:
        (rep / f"{r['gold_id']}_{r['case_id']}.txt").write_text(load_text(index[r["case_id"]]))
    (OUT / "sample_key.json").write_text(json.dumps(
        {"seed": SEED, "n_random": N_RANDOM,
         "enrich": [(w, k) for w, _, k in ENRICH],
         "note": "Do not open before labelling is finished - it names each case's stratum.",
         "rows": rows}, indent=1))
    (OUT / "GUIDE.md").write_text("```\n" + GUIDE.strip("\n") + "\n```\n")
    build_workbook(rows, "Labiba").save(OUT / "gold_labels_A.xlsx")
    build_workbook(rows, "B").save(OUT / "gold_labels_B.xlsx")
    build_practice(index)
    (OUT / "CHEAT_SHEET.md").write_text("```\n" + CHEAT.strip("\n") + "\n```\n")
    print(f"wrote {OUT.relative_to(BASE)}/: gold_labels_A.xlsx, gold_labels_B.xlsx, "
          f"GUIDE.md, sample_key.json, reports/ ({len(rows)} files)")


if __name__ == "__main__":
    build()
