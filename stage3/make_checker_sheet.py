"""make_checker_sheet.py - improvement item 3: a blind reading sheet for 60 LLM-written reports.

THESIS_LOG.md, 29 Sep (plan fixed before building). Writes, under checker_validation/:
  GUIDE.md                     how to read and fill the sheet (also the first tab)
  checker_sheet_A.xlsx         Labiba's sheet: report text + dropdowns, nothing else
  reports/C01.txt ...          the report texts (input for the LLM judge)
  key/key.json                 DO NOT OPEN before the sheet is done: model, arm, case, and the
                               checker's own reading (v1.3), frozen here before any reader is seen
"""
from __future__ import annotations

import hashlib
import json
import random
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
BASE = HERE.parent
sys.path.insert(0, str(HERE))
import consistency as C   # noqa: E402

OUT = BASE / "checker_validation"
MODELS = ["medgemma-27b-text-it", "Llama-3.1-8B-Instruct", "medgemma-1.5-4b-it"]
ARMS = ["predicted", "extracted", "none_free", "retrieval"]
PER_CELL = 5
SEED = 29
EXTRA = [("lymphovascular_invasion", "LVI"), ("margins", "Margins"), ("node_count", "Node count"),
         ("pn_category", "pN"), ("tumour_size", "Size"), ("distant_metastasis", "Distant met / M"),
         ("msi_mmr", "MSI/MMR"), ("perineural_invasion", "Perineural"), ("tumour_budding", "Budding"),
         ("other_staging", "Other staging"), ("treatment_history", "Treatment"), ("tumour_site", "Site")]
# "Site" added 29 Sep before any human reading: the checker also flags tumour site (found elsewhere in
# consistency.py than the UNSUPPORTED table the first list was copied from)
OPTIONS = {
    "Type": ["not stated", "adenocarcinoma", "mucinous", "signet-ring"],
    "Grade": ["not stated", "G1", "G2", "G3", "G4"],
    "pT": ["not stated", "pT1", "pT2", "pT3", "pT4"],
    "Node": ["not stated", "cannot be assessed", "negative", "positive"],
    "Node hedged": ["yes", "no"],
    "yes": ["yes"],
}

GUIDE = """# Reading guide: what does each generated report SAY?

60 short reports written by language models (C01 to C60). For each one, record **what the report
states**, not whether it is true. You do not know the patient, the model or its input, on purpose.
Takes about 45 to 60 minutes. Fill the "Labels" tab; use the dropdowns.

## The four findings

| Column | Choose | Notes |
|---|---|---|
| Type | not stated / adenocarcinoma / mucinous / signet-ring | "Mucinous adenocarcinoma" = mucinous. "Adenocarcinoma with a mucinous component" or "focal mucinous features" = adenocarcinoma (WHO: mucinous means more than 50% mucin). "Not mucinous" = adenocarcinoma. A confidence number such as "(model confidence 40%)" is NOT a mucin percentage. |
| Grade | not stated / G1 to G4 | well = G1, moderately = G2, poorly = G3, undifferentiated = G4. "Low grade" alone = not stated (it merges G1 and G2). "Grade cannot be determined" = not stated. |
| pT | not stated / pT1 to pT4 | Words count: "invades the muscularis propria" = pT2; "through the muscularis propria into pericolorectal/subserosal tissue" = pT3; "submucosa" = pT1; "visceral peritoneum" or "adjacent organ" = pT4. pT4a/pT4b = pT4. |
| Node | not stated / cannot be assessed / negative / positive | "Cannot be assessed" = the report says node status cannot be determined from this tissue. |
| Node hedged | yes / no | Only if Node is negative or positive: yes if the node statement is softened ("likely", "estimated", "predicted", "may", "model confidence"). Leave blank otherwise. |

If a report states two different values for one finding, pick the one stated as the diagnosis and
write the conflict in Notes.

## Extra claims (anything beyond the four findings)

Put **yes** in a column if the report makes that claim at all, positive or negative
("margins are free" is a margin claim). Leave blank if not.

| Column | Examples |
|---|---|
| LVI | lymphovascular, venous, vascular invasion; L0/L1, V0/V1 |
| Margins | margins free/involved, R0/R1 |
| Node count | "0 of 15 lymph nodes", "3/12 nodes" |
| pN | pN0, pN1, N2 |
| Size | "4.5 cm", "30 mm" |
| Distant met / M | pM1, M0, liver metastasis |
| MSI/MMR | MSI, MSS, mismatch repair, MLH1 |
| Perineural | perineural invasion |
| Budding | tumour budding |
| Other staging | Dukes, AJCC stage, stage II |
| Treatment | neoadjuvant, treatment effect, regression grade |
| Site | where in the bowel the tumour is: caecum, ascending, transverse, descending or sigmoid colon, hepatic/splenic flexure, rectosigmoid, rectum. "Colon" or "colorectal" alone does not count. |

**Other claim**: free text for any other factual claim about THIS patient's tissue that none of the
columns covers (for example necrosis, a polyp, ulceration). Leave blank if none.

## Rules

- Record what is written, even if it looks wrong or contradicts itself.
- Do not open `checker_validation/key/` until your sheet is finished.
- When unsure, pick your best answer and write why in Notes.
"""


def checker_reading(row):
    chk = C.check(row["output"], row["input_findings"])
    fl = chk["fields"]
    node = fl["node_status"]["stated"]
    node_abst = any(C.NODE_WORDS.search(s) for s in chk["abstentions"])
    hedged = C.extract_claims(row["output"])[1].get("node_status") if node is not None else None
    t = fl["histologic_type"]["stated"]
    return {"Type": {"adenocarcinoma": "adenocarcinoma", "mucinous": "mucinous",
                     "signet_ring": "signet-ring", None: "not stated"}.get(t, str(t)),
            "Grade": f"G{fl['grade']['stated']}" if fl["grade"]["stated"] else "not stated",
            "pT": f"pT{fl['pt_stage']['stated']}" if fl["pt_stage"]["stated"] else "not stated",
            "Node": ("positive" if node else "negative") if node is not None
                    else ("cannot be assessed" if node_abst else "not stated"),
            "Node hedged": None if node is None else ("yes" if hedged else "no"),
            "extra": sorted(chk["unsupported_claims"].keys())}


def main():
    gold = {r["case_id"] for r in json.loads((BASE / "gold_labels" / "sample_key.json").read_text())["rows"]}
    rng = random.Random(SEED)
    used, picks = set(), []
    for m in MODELS:
        for arm in ARMS:
            rows = [json.loads(l) for l in (BASE / "stage2" / "outputs" / f"{m}_tcga" / f"{arm}.jsonl").read_text().splitlines() if l.strip()]
            pool = [r for r in rows if r["case_id"] not in gold and r["case_id"] not in used]
            for r in rng.sample(pool, PER_CELL):
                used.add(r["case_id"]); picks.append((m, arm, r))
    rng.shuffle(picks)
    (OUT / "reports").mkdir(parents=True, exist_ok=True); (OUT / "key").mkdir(exist_ok=True)
    key = []
    for i, (m, arm, r) in enumerate(picks, 1):
        rid = f"C{i:02d}"
        (OUT / "reports" / f"{rid}.txt").write_text(r["output"].strip() + "\n")
        key.append({"id": rid, "case_id": r["case_id"], "split": r["split"], "model": m, "arm": arm,
                    "input_findings": r["input_findings"], "reference_findings": r["reference_findings"],
                    "checker_version": C.VERSION, "checker": checker_reading(r),
                    "text_sha256": hashlib.sha256(r["output"].strip().encode()).hexdigest()})
    (OUT / "key" / "key.json").write_text(json.dumps({"seed": SEED, "plan": "THESIS_LOG.md 29 Sep item 3",
                                                      "rows": key}, indent=1))
    (OUT / "GUIDE.md").write_text(GUIDE)

    wb = Workbook()
    g = wb.active; g.title = "Guide"
    for line in GUIDE.splitlines():
        g.append([line])
    g.column_dimensions["A"].width = 140
    ws = wb.create_sheet("Labels")
    cols = ["ID", "Report", "Type", "Grade", "pT", "Node", "Node hedged"] + [lab for _, lab in EXTRA] + ["Other claim", "Notes"]
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDEBF7")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for k in key:
        ws.append([k["id"], (OUT / "reports" / f"{k['id']}.txt").read_text().strip()] + [None] * (len(cols) - 2))
    n = len(key) + 1
    ws.column_dimensions["A"].width = 6; ws.column_dimensions["B"].width = 80
    for j in range(3, len(cols) + 1):
        ws.column_dimensions[ws.cell(1, j).column_letter].width = 13
    for row in ws.iter_rows(min_row=2, max_row=n):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    lists = wb.create_sheet("Lists")
    for j, (name, opts) in enumerate(OPTIONS.items(), 1):
        for i, o in enumerate(opts, 1):
            lists.cell(i, j, o)
    lists.sheet_state = "hidden"
    ref = {name: f"Lists!${lists.cell(1, j).column_letter}$1:${lists.cell(1, j).column_letter}${len(opts)}"
           for j, (name, opts) in enumerate(OPTIONS.items(), 1)}
    for j, col in enumerate(cols, 1):
        key_ = col if col in OPTIONS else ("yes" if col in [lab for _, lab in EXTRA] else None)
        if key_:
            dv = DataValidation(type="list", formula1=ref[key_], allow_blank=True)
            L = ws.cell(1, j).column_letter
            dv.add(f"{L}2:{L}{n}"); ws.add_data_validation(dv)
    ws.freeze_panes = "C2"
    wb.save(OUT / "checker_sheet_A.xlsx")
    from collections import Counter
    print(len(key), "reports;", Counter((k["model"], k["arm"]) for k in key).most_common(2), "...")
    print("wrote checker_validation/{GUIDE.md, checker_sheet_A.xlsx, reports/, key/key.json}")


if __name__ == "__main__":
    main()
