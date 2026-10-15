"""score_checker_validation.py - improvement item 3: checker vs independent readers on 60 LLM reports.

THESIS_LOG.md, 29 Sep (plan fixed before building). Readers:
  checker  checker_validation/key/key.json (checker v1.3 reading, frozen when the sheet was built)
  judge    checker_validation/llm_judge/judge_claude_raw.json (Claude, blind; SHA-256 in the log)
  human    checker_validation/checker_sheet_A.xlsx, once filled (skipped while empty)
Per field: % agreement and Cohen's kappa for every pair. Extra claims: report-level presence and
per-kind agreement. Every disagreement listed for adjudication.
Writes stage3/checker_validation_results.json.
"""
from __future__ import annotations

import json
from collections import Counter
from pathlib import Path

HERE = Path(__file__).resolve().parent
BASE = HERE.parent
CV = BASE / "checker_validation"
FIELDS = ["Type", "Grade", "pT", "Node", "Node hedged"]
KINDS = {"lymphovascular_invasion": "LVI", "margins": "Margins", "node_count": "Node count", "pn_category": "pN",
         "tumour_size": "Size", "distant_metastasis": "Distant met / M", "msi_mmr": "MSI/MMR",
         "perineural_invasion": "Perineural", "tumour_budding": "Budding", "other_staging": "Other staging",
         "treatment_history": "Treatment", "tumour_site": "Site"}


def kappa(a, b):
    n = len(a)
    if n == 0:
        return None
    po = sum(x == y for x, y in zip(a, b)) / n
    ca, cb = Counter(a), Counter(b)
    pe = sum(ca[k] * cb[k] for k in set(a) | set(b)) / n ** 2
    return None if pe == 1 else (po - pe) / (1 - pe)


def load():
    key = json.loads((CV / "key" / "key.json").read_text())["rows"]
    R = {"checker": {k["id"]: {**{f: k["checker"][f] for f in FIELDS},
                               "extra": sorted(KINDS[x] for x in k["checker"]["extra"])} for k in key}}
    # the CURRENT checker re-read (secondary; the frozen v1.3 reading above is the primary result)
    import sys
    sys.path.insert(0, str(HERE))
    import consistency as C
    import make_checker_sheet as M
    if C.VERSION != key[0]["checker_version"]:
        rows = {}
        for m in M.MODELS:
            for arm in M.ARMS:
                for l in (BASE / "stage2" / "outputs" / f"{m}_tcga" / f"{arm}.jsonl").read_text().splitlines():
                    r = json.loads(l); rows[(m, arm, r["case_id"])] = r
        cur = {k["id"]: M.checker_reading(rows[(k["model"], k["arm"], k["case_id"])]) for k in key}
        R[f"checker_{C.VERSION}"] = {i: {**{f: v[f] for f in FIELDS}, "extra": sorted(KINDS[x] for x in v["extra"])}
                                     for i, v in cur.items()}
    j = CV / "llm_judge" / "judge_claude_raw.json"
    if j.exists():
        R["judge"] = {r["id"]: {**{f: r[f] for f in FIELDS}, "extra": sorted(r["extra"]),
                                "other_claim": r["other_claim"], "notes": r["notes"]} for r in json.loads(j.read_text())}
        site = CV / "llm_judge" / "judge_claude_site.json"     # Site asked separately (column added later)
        if site.exists():
            for i, v in json.loads(site.read_text()).items():
                if v == "yes":
                    R["judge"][i]["extra"] = sorted(R["judge"][i]["extra"] + ["Site"])
    x = CV / "checker_sheet_A.xlsx"
    if x.exists():
        from openpyxl import load_workbook
        ws = load_workbook(x, data_only=True)["Labels"]
        head = [c.value for c in ws[1]]
        rows = {}
        for r in ws.iter_rows(min_row=2, values_only=True):
            d = dict(zip(head, r))
            if not any(d.get(f) for f in ["Type", "Grade", "pT", "Node"]):
                continue
            rows[d["ID"]] = {**{f: d.get(f) for f in FIELDS},
                             "extra": sorted(lab for lab in KINDS.values() if (d.get(lab) or "").strip().lower() == "yes"),
                             "other_claim": d.get("Other claim") or "", "notes": d.get("Notes") or ""}
        if rows:
            R["human"] = rows
    return key, R


def compare(A, B, ids):
    out = {}
    for f in FIELDS:
        pairs = [(A[i][f], B[i][f]) for i in ids]
        if f == "Node hedged":      # only where both readers say node is stated
            pairs = [(A[i][f], B[i][f]) for i in ids if A[i]["Node"] in ("negative", "positive")
                     and B[i]["Node"] in ("negative", "positive")]
        a, b = [str(p[0]) for p in pairs], [str(p[1]) for p in pairs]
        out[f] = {"n": len(pairs), "agreement": sum(x == y for x, y in zip(a, b)) / len(pairs) if pairs else None,
                  "kappa": kappa(a, b)}
    anyA, anyB = [bool(A[i]["extra"]) for i in ids], [bool(B[i]["extra"]) for i in ids]
    out["any extra claim"] = {"n": len(ids), "agreement": sum(x == y for x, y in zip(anyA, anyB)) / len(ids),
                              "kappa": kappa(anyA, anyB)}
    per = {}
    for lab in KINDS.values():
        both = sum(lab in A[i]["extra"] and lab in B[i]["extra"] for i in ids)
        a_only = sum(lab in A[i]["extra"] and lab not in B[i]["extra"] for i in ids)
        b_only = sum(lab not in A[i]["extra"] and lab in B[i]["extra"] for i in ids)
        if both + a_only + b_only:
            per[lab] = {"both": both, "first_only": a_only, "second_only": b_only}
    out["extra_by_kind"] = per
    return out


def main():
    key, R = load()
    ids = [k["id"] for k in key]
    names = list(R)
    res = {"readers": names, "n_reports": len(ids),
           "n_distinct_texts": len({k["text_sha256"] for k in key}), "pairs": {}, "disagreements": {}}
    for i, a in enumerate(names):
        for b in names[i + 1:]:
            common = [x for x in ids if x in R[a] and x in R[b]]
            res["pairs"][f"{a} vs {b}"] = compare(R[a], R[b], common)
            dis = []
            for x in common:
                d = {f: (R[a][x][f], R[b][x][f]) for f in FIELDS if str(R[a][x][f]) != str(R[b][x][f])}
                if R[a][x]["extra"] != R[b][x]["extra"]:
                    d["extra"] = (R[a][x]["extra"], R[b][x]["extra"])
                if d:
                    dis.append({"id": x, **{k: list(v) for k, v in d.items()}})
            res["disagreements"][f"{a} vs {b}"] = dis
    (HERE / "checker_validation_results.json").write_text(json.dumps(res, indent=1))
    f = lambda v: "  -  " if v is None else f"{v:.2f}"
    print(f"{len(ids)} reports ({res['n_distinct_texts']} distinct texts); readers: {', '.join(names)}")
    for p, c in res["pairs"].items():
        print(f"\n== {p}")
        for k in FIELDS + ["any extra claim"]:
            print(f"  {k:16s} n={c[k]['n']:3d}  agreement {f(c[k]['agreement'])}  kappa {f(c[k]['kappa'])}")
        print("  extra by kind (both / first only / second only):",
              "; ".join(f"{k} {v['both']}/{v['first_only']}/{v['second_only']}" for k, v in c["extra_by_kind"].items()))
        print(f"  disagreements: {len(res['disagreements'][p])} reports")


if __name__ == "__main__":
    main()
